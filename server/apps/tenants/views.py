import csv
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO, StringIO

from django.contrib.auth import get_user_model
from django.core.mail import send_mail
from django.db import transaction
from django.db.models import Count, Sum
from django.db.models.functions import Coalesce
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework import status, viewsets
from rest_framework.decorators import action
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.accounts.mixins import OrgScopedMixin
from apps.crm.models import ContactAccount
from apps.facility_management.models import Facility, FacilityUnitSpace
from apps.finance.models import Customer, Invoice, InvoiceLineItem, InvoicePayment
from apps.properties.models import Property, Unit, WorkOrder
from apps.settings.permissions import HasRolePermission
from apps.support_desk.whatsapp_providers import get_whatsapp_provider_adapter

from .models import (
    TenantCommunicationLog,
    TenantIdentityDocument,
    TenantIncidentRecord,
    TenantInventoryItem,
    TenantProfile,
    TenantRelationshipContact,
)
from .serializers import (
    TenantCommunicationLogSerializer,
    TenantIdentityDocumentSerializer,
    TenantIdentityDocumentWriteSerializer,
    TenantIncidentRecordSerializer,
    TenantIncidentRecordWriteSerializer,
    TenantInventoryItemSerializer,
    TenantInventoryItemWriteSerializer,
    TenantLedgerLineItemSerializer,
    TenantLedgerPaymentSerializer,
    TenantProfileDetailSerializer,
    TenantProfileListSerializer,
    TenantProfileWriteSerializer,
    TenantRelationshipContactSerializer,
    TenantRelationshipContactWriteSerializer,
)

NOTICE_HORIZON_DAYS = 45
ZERO_DECIMAL = Decimal("0.00")
User = get_user_model()

TENANT_REGISTRY_IMPORT_EXPORT_COLUMNS = [
    "display_name",
    "customer_name",
    "customer_email",
    "customer_phone",
    "contact_account_name",
    "contact_email",
    "contact_phone",
    "primary_user_email",
    "property_name",
    "property_location",
    "unit_number",
    "facility_code",
    "space_label",
    "tenant_type",
    "status",
    "lease_start_date",
    "lease_end_date",
    "move_in_date",
    "move_out_date",
    "occupant_count",
    "notes",
]

TENANT_TYPE_IMPORT_ALIASES = {
    "individual": "individual",
    "residential": "individual",
    "corporate": "corporate",
    "commercial": "corporate",
}

TENANT_STATUS_IMPORT_ALIASES = {
    "active": "active",
    "prospective": "pending_move_in",
    "pending move in": "pending_move_in",
    "pending_move_in": "pending_move_in",
    "past": "moved_out",
    "archived": "moved_out",
    "moved out": "moved_out",
    "moved_out": "moved_out",
    "inactive": "inactive",
}


def _resolve_user_org(request):
    org = getattr(request, "organization", None)
    if org is not None:
        return org
    profile = getattr(getattr(request, "user", None), "profile", None)
    return getattr(profile, "organization", None)


def _month_bounds(value):
    month_start = value.replace(day=1)
    if month_start.month == 12:
        next_month_start = month_start.replace(year=month_start.year + 1, month=1)
    else:
        next_month_start = month_start.replace(month=month_start.month + 1)
    return month_start, next_month_start


def _add_months(value, months):
    month_index = (value.month - 1) + months
    year = value.year + (month_index // 12)
    month = (month_index % 12) + 1
    return value.replace(year=year, month=month, day=1)


def _month_label(value) -> str:
    return value.strftime("%b %Y")


def _resolve_profile_property(profile):
    if profile.property_id and profile.property:
        return profile.property
    if profile.unit_id and profile.unit:
        return profile.unit.property
    if profile.facility_id and profile.facility:
        return profile.facility.property
    if profile.facility_space_id and profile.facility_space:
        return profile.facility_space.facility.property
    return None


def _resolve_profile_property_name(profile) -> str:
    property_record = _resolve_profile_property(profile)
    return property_record.name if property_record else "--"


def _property_location_label(property_record) -> str:
    if not property_record or not property_record.address:
        return "Unknown"
    parts = [part.strip() for part in property_record.address.split(",") if part.strip()]
    return parts[-1] if parts else "Unknown"


def _resolve_profile_property_location(profile) -> str:
    return _property_location_label(_resolve_profile_property(profile))


def _resolve_profile_unit_label(profile) -> str:
    if profile.unit_id and profile.unit:
        return profile.unit.unit_number
    if profile.facility_space_id and profile.facility_space and profile.facility_space.unit_id:
        return profile.facility_space.unit.unit_number
    if profile.facility_space_id and profile.facility_space and profile.facility_space.space_label:
        return profile.facility_space.space_label
    return "--"


def _resolve_profile_unit_category(profile) -> str:
    if profile.unit_id and profile.unit:
        return profile.unit.unit_category
    if profile.facility_space_id and profile.facility_space and profile.facility_space.unit_id:
        return profile.facility_space.unit.unit_category
    return ""


def _resolve_profile_tenancy_type(profile):
    residential_categories = {
        Unit.UnitCategory.APARTMENT,
        Unit.UnitCategory.VILLA,
        Unit.UnitCategory.TOWNHOUSE,
        Unit.UnitCategory.PENTHOUSE,
        Unit.UnitCategory.STUDIO,
        Unit.UnitCategory.DUPLEX,
    }
    unit_category = _resolve_profile_unit_category(profile)
    if unit_category in residential_categories:
        return "residential", "Residential"
    if unit_category:
        return "commercial", "Commercial"
    property_record = _resolve_profile_property(profile)
    if property_record and property_record.property_type in {
        Property.PropertyType.WAREHOUSE,
        Property.PropertyType.INDUSTRIAL,
    }:
        return "commercial", "Commercial"
    return "residential", "Residential"


def _profile_notice_date(profile):
    return profile.move_out_date or profile.lease_end_date


def _is_notice_given(profile, today, *, horizon_days=NOTICE_HORIZON_DAYS) -> bool:
    if profile.status != TenantProfile.Status.ACTIVE:
        return False
    notice_date = _profile_notice_date(profile)
    if not notice_date:
        return False
    return today <= notice_date <= today + timedelta(days=horizon_days)


def _tenant_dashboard_status(profile, customer_snapshot, today):
    if customer_snapshot.get("overdue_invoices", 0):
        return "overdue", "Overdue"
    if profile.status == TenantProfile.Status.ACTIVE and _is_notice_given(profile, today):
        return "notice_given", "Notice Given"
    if profile.status == TenantProfile.Status.ACTIVE:
        return "active", "Active"
    if profile.status == TenantProfile.Status.PENDING_MOVE_IN:
        return "pending_move_in", "Pending Move In"
    if profile.status == TenantProfile.Status.MOVED_OUT:
        return "moved_out", "Moved Out"
    return "inactive", "Inactive"


def _tenant_debt_status(customer_snapshot):
    balance_due = customer_snapshot.get("balance_due", ZERO_DECIMAL) or ZERO_DECIMAL
    overdue_invoices = customer_snapshot.get("overdue_invoices", 0) or 0
    if overdue_invoices:
        return "overdue", "Overdue"
    if balance_due > ZERO_DECIMAL:
        return "outstanding", "Outstanding"
    return "clear", "Clear"


def _announcement_subject(subject: str) -> str:
    cleaned = (subject or "").strip()
    return cleaned or "Holiday announcement"


def _normalize_lookup_value(value) -> str:
    return " ".join(str(value or "").strip().lower().split())


def _normalize_registry_cell(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value or "").strip()


def _registry_export_row(profile) -> dict[str, str]:
    return {
        "display_name": profile.resolved_display_name,
        "customer_name": profile.customer.name if profile.customer_id and profile.customer else "",
        "customer_email": profile.customer.email if profile.customer_id and profile.customer else "",
        "customer_phone": profile.customer.phone if profile.customer_id and profile.customer else "",
        "contact_account_name": profile.contact_account.display_name if profile.contact_account_id and profile.contact_account else "",
        "contact_email": _tenant_contact_email(profile),
        "contact_phone": _tenant_contact_phone(profile),
        "primary_user_email": profile.primary_user.email if profile.primary_user_id and profile.primary_user else "",
        "property_name": _resolve_profile_property_name(profile),
        "property_location": _resolve_profile_property_location(profile),
        "unit_number": _resolve_profile_unit_label(profile),
        "facility_code": profile.facility.facility_code if profile.facility_id and profile.facility else (
            profile.facility_space.facility.facility_code if profile.facility_space_id and profile.facility_space else ""
        ),
        "space_label": profile.facility_space.space_label if profile.facility_space_id and profile.facility_space else "",
        "tenant_type": profile.get_tenant_type_display(),
        "status": profile.get_status_display(),
        "lease_start_date": profile.lease_start_date.isoformat() if profile.lease_start_date else "",
        "lease_end_date": profile.lease_end_date.isoformat() if profile.lease_end_date else "",
        "move_in_date": profile.move_in_date.isoformat() if profile.move_in_date else "",
        "move_out_date": profile.move_out_date.isoformat() if profile.move_out_date else "",
        "occupant_count": str(profile.occupant_count or 1),
        "notes": profile.notes or "",
    }


def _load_registry_import_rows(upload):
    file_name = (getattr(upload, "name", "") or "").lower()
    if file_name.endswith(".csv"):
        content = upload.read().decode("utf-8-sig")
        reader = csv.DictReader(StringIO(content))
        return [{key: _normalize_registry_cell(value) for key, value in row.items()} for row in reader], "csv"

    if file_name.endswith(".xlsx"):
        workbook = load_workbook(filename=BytesIO(upload.read()), data_only=True, read_only=True)
        worksheet = workbook.active
        raw_rows = list(worksheet.iter_rows(values_only=True))
        if not raw_rows:
            return [], "xlsx"
        headers = [str(cell or "").strip() for cell in raw_rows[0]]
        rows = []
        for row_values in raw_rows[1:]:
            if row_values is None:
                continue
            row_dict = {
                headers[index]: _normalize_registry_cell(row_values[index]) if index < len(row_values) else ""
                for index in range(len(headers))
            }
            if any(str(value or "").strip() for value in row_dict.values()):
                rows.append(row_dict)
        return rows, "xlsx"

    raise ValueError("Only CSV and XLSX files are supported.")


def _tenant_registry_export_response(*, queryset, format_key):
    if format_key not in {"csv", "xlsx"}:
        return Response({"detail": "format must be csv or xlsx."}, status=status.HTTP_400_BAD_REQUEST)

    rows = [_registry_export_row(profile) for profile in queryset.order_by("display_name", "id")]
    date_stamp = timezone.localdate().isoformat()

    if format_key == "csv":
        response = HttpResponse(content_type="text/csv")
        response["Content-Disposition"] = f'attachment; filename="tenant-registry-{date_stamp}.csv"'
        writer = csv.DictWriter(response, fieldnames=TENANT_REGISTRY_IMPORT_EXPORT_COLUMNS)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)
        return response

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Tenant Registry"
    worksheet.append(TENANT_REGISTRY_IMPORT_EXPORT_COLUMNS)
    for row in rows:
        worksheet.append([row.get(column, "") for column in TENANT_REGISTRY_IMPORT_EXPORT_COLUMNS])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="tenant-registry-{date_stamp}.xlsx"'
    return response


def _parse_positive_decimal(value, *, field_name):
    try:
        amount = Decimal(str(value))
    except Exception as exc:  # pragma: no cover - defensive parsing
        raise ValueError(f"{field_name} must be a valid amount.") from exc
    if amount <= ZERO_DECIMAL:
        raise ValueError(f"{field_name} must be greater than zero.")
    return amount.quantize(Decimal("0.01"))


def _invoice_snapshot(*, organization, customer_ids, today, month_start=None, next_month_start=None):
    snapshot = {
        "by_customer": {},
        "totals": {
            "overdue_total": ZERO_DECIMAL,
            "overdue_invoice_count": 0,
            "customers_in_arrears": set(),
            "current_month_target_total": ZERO_DECIMAL,
            "current_month_collections": ZERO_DECIMAL,
        },
    }
    if not customer_ids:
        return snapshot

    invoice_queryset = (
        Invoice.objects.filter(organization=organization, customer_id__in=customer_ids)
        .exclude(status=Invoice.Status.CANCELLED)
        .prefetch_related("payments")
        .order_by("-due_date", "-id")
    )

    for invoice in invoice_queryset:
        paid_amount = sum((payment.amount for payment in invoice.payments.all()), ZERO_DECIMAL)
        balance_due = max(invoice.total_amount - paid_amount, ZERO_DECIMAL)
        customer_summary = snapshot["by_customer"].setdefault(
            invoice.customer_id,
            {
                "balance_due": ZERO_DECIMAL,
                "overdue_balance": ZERO_DECIMAL,
                "overdue_invoices": 0,
            },
        )
        customer_summary["balance_due"] += balance_due

        if balance_due > 0 and invoice.due_date < today:
            customer_summary["overdue_balance"] += balance_due
            customer_summary["overdue_invoices"] += 1
            snapshot["totals"]["overdue_total"] += balance_due
            snapshot["totals"]["overdue_invoice_count"] += 1
            snapshot["totals"]["customers_in_arrears"].add(invoice.customer_id)

        if (
            month_start
            and next_month_start
            and invoice.status != Invoice.Status.DRAFT
            and month_start <= invoice.due_date < next_month_start
        ):
            snapshot["totals"]["current_month_target_total"] += invoice.total_amount

    if month_start and next_month_start:
        snapshot["totals"]["current_month_collections"] = (
            InvoicePayment.objects.filter(
                invoice__organization=organization,
                invoice__customer_id__in=customer_ids,
                payment_date__gte=month_start,
                payment_date__lt=next_month_start,
            ).aggregate(total=Coalesce(Sum("amount"), ZERO_DECIMAL))["total"]
            or ZERO_DECIMAL
        )

    return snapshot


def _tenant_customer_ids_for_org(organization):
    return sorted(
        {
            customer_id
            for customer_id in TenantProfile.objects.filter(organization=organization).values_list("customer_id", flat=True)
            if customer_id
        }
    )


def _open_work_orders_by_unit(organization):
    return {
        row["unit_id"]: row["count"]
        for row in WorkOrder.objects.filter(organization=organization)
        .exclude(
            status__in=[
                WorkOrder.Status.COMPLETED,
                WorkOrder.Status.VERIFIED,
                WorkOrder.Status.CANCELLED,
            ]
        )
        .filter(unit_id__isnull=False)
        .values("unit_id")
        .annotate(count=Count("id"))
    }


def _tenant_invoice_queryset(profile):
    if not profile.customer_id:
        return Invoice.objects.none()
    return (
        Invoice.objects.filter(organization=profile.organization, customer_id=profile.customer_id)
        .exclude(status=Invoice.Status.CANCELLED)
        .prefetch_related("payments")
        .order_by("-issue_date", "-due_date", "-id")
    )


def _tenant_contact_email(profile) -> str:
    if profile.primary_user_id and profile.primary_user and profile.primary_user.email:
        return profile.primary_user.email.strip()
    if profile.customer_id and profile.customer and profile.customer.email:
        return profile.customer.email.strip()
    if profile.contact_account_id and profile.contact_account and profile.contact_account.email:
        return profile.contact_account.email.strip()
    return ""


def _tenant_contact_phone(profile) -> str:
    if profile.customer_id and profile.customer and profile.customer.phone:
        return profile.customer.phone.strip()
    if profile.contact_account_id and profile.contact_account:
        if profile.contact_account.phone:
            return profile.contact_account.phone.strip()
        if profile.contact_account.secondary_phone:
            return profile.contact_account.secondary_phone.strip()
    return ""


def _tenant_invoice_metrics(profile, today):
    metrics = {
        "total_invoices": 0,
        "paid_on_time": 0,
        "paid_late": 0,
        "overdue_invoices": 0,
        "current_balance": ZERO_DECIMAL,
        "overdue_balance": ZERO_DECIMAL,
        "latest_invoice_amount": ZERO_DECIMAL,
        "latest_invoice_number": "",
        "latest_due_date": None,
        "oldest_overdue_due_date": None,
    }

    invoices = list(_tenant_invoice_queryset(profile))
    if invoices:
        metrics["latest_invoice_amount"] = invoices[0].total_amount or ZERO_DECIMAL
        metrics["latest_invoice_number"] = invoices[0].invoice_number or ""
        metrics["latest_due_date"] = invoices[0].due_date

    for invoice in invoices:
        metrics["total_invoices"] += 1
        payments = list(invoice.payments.all())
        paid_amount = sum((payment.amount for payment in payments), ZERO_DECIMAL)
        balance_due = max((invoice.total_amount or ZERO_DECIMAL) - paid_amount, ZERO_DECIMAL)
        metrics["current_balance"] += balance_due

        if balance_due > ZERO_DECIMAL and invoice.due_date < today:
            metrics["overdue_invoices"] += 1
            metrics["overdue_balance"] += balance_due
            if metrics["oldest_overdue_due_date"] is None or invoice.due_date < metrics["oldest_overdue_due_date"]:
                metrics["oldest_overdue_due_date"] = invoice.due_date

        if invoice.total_amount > ZERO_DECIMAL and paid_amount >= invoice.total_amount:
            last_payment_date = max((payment.payment_date for payment in payments), default=None)
            if last_payment_date and last_payment_date <= invoice.due_date:
                metrics["paid_on_time"] += 1
            elif last_payment_date:
                metrics["paid_late"] += 1

    return metrics


def _payment_reliability_snapshot(metrics):
    total_invoices = metrics["total_invoices"]
    if total_invoices == 0:
        return {
            "score": 3,
            "label": "Neutral",
            "summary": "No receivables history yet, so the tenant is shown with a neutral baseline score.",
            "total_invoices": 0,
            "paid_on_time": 0,
            "paid_late": 0,
            "overdue_invoices": 0,
            "current_balance": format(metrics["current_balance"], ".2f"),
            "overdue_balance": format(metrics["overdue_balance"], ".2f"),
        }

    late_ratio = metrics["paid_late"] / total_invoices
    overdue_ratio = metrics["overdue_invoices"] / total_invoices
    score = max(1, min(5, int((5 - (late_ratio * 2.0) - (overdue_ratio * 3.0)) + 0.5)))

    label_map = {
        5: "Excellent",
        4: "Strong",
        3: "Fair",
        2: "Weak",
        1: "Critical",
    }
    if score >= 5:
        summary = "Pays on time consistently with no overdue exposure across recent invoices."
    elif score == 4:
        summary = "Generally reliable, with only minor slippage in payment cadence."
    elif score == 3:
        summary = "Mixed payment behavior. Monitoring and proactive reminders are advised."
    elif score == 2:
        summary = "Collection risk is elevated due to late payments or active arrears."
    else:
        summary = "Severe payment risk. Immediate recovery action is recommended."

    return {
        "score": score,
        "label": label_map[score],
        "summary": summary,
        "total_invoices": total_invoices,
        "paid_on_time": metrics["paid_on_time"],
        "paid_late": metrics["paid_late"],
        "overdue_invoices": metrics["overdue_invoices"],
        "current_balance": format(metrics["current_balance"], ".2f"),
        "overdue_balance": format(metrics["overdue_balance"], ".2f"),
    }


def _lease_timeline_snapshot(profile, today):
    start_date = profile.lease_start_date or profile.move_in_date
    end_date = profile.move_out_date or profile.lease_end_date
    if not start_date or not end_date or end_date <= start_date:
        return {
            "start_date": start_date.isoformat() if start_date else None,
            "end_date": end_date.isoformat() if end_date else None,
            "elapsed_days": 0,
            "remaining_days": 0,
            "total_days": 0,
            "elapsed_pct": 0.0,
            "status_label": "Unavailable",
        }

    total_days = max((end_date - start_date).days, 1)
    if today <= start_date:
        elapsed_days = 0
        remaining_days = total_days
        status_label = "Upcoming"
    elif today >= end_date:
        elapsed_days = total_days
        remaining_days = 0
        status_label = "Completed"
    else:
        elapsed_days = (today - start_date).days
        remaining_days = (end_date - today).days
        status_label = "Active"

    return {
        "start_date": start_date.isoformat(),
        "end_date": end_date.isoformat(),
        "elapsed_days": elapsed_days,
        "remaining_days": remaining_days,
        "total_days": total_days,
        "elapsed_pct": round((elapsed_days / total_days) * 100, 1),
        "status_label": status_label,
    }


def _record_tenant_log(
    profile,
    *,
    interaction_type,
    channel,
    direction,
    entry_status,
    subject,
    message,
    author=None,
    metadata=None,
):
    return TenantCommunicationLog.objects.create(
        organization=profile.organization,
        tenant_profile=profile,
        author=author,
        interaction_type=interaction_type,
        channel=channel,
        direction=direction,
        status=entry_status,
        subject=subject,
        message=message,
        metadata=metadata or {},
    )


def _serialize_communication_log_entries(profile):
    entries = list(profile.communication_logs.select_related("author").all()[:8])
    payload = list(TenantCommunicationLogSerializer(entries, many=True).data)
    if profile.notes:
        payload.append(
            {
                "id": f"profile-note-{profile.id}",
                "interaction_type": TenantCommunicationLog.InteractionType.NOTE,
                "interaction_type_display": "Profile Note",
                "channel": TenantCommunicationLog.Channel.INTERNAL_NOTE,
                "channel_display": "Internal Note",
                "direction": TenantCommunicationLog.Direction.INTERNAL,
                "direction_display": "Internal",
                "status": TenantCommunicationLog.Status.RECORDED,
                "status_display": "Recorded",
                "subject": "Profile note",
                "message": profile.notes,
                "metadata": {"source": "tenant_profile"},
                "happened_at": profile.updated_at.isoformat(),
                "author_name": "",
            }
        )
    payload.sort(key=lambda item: item.get("happened_at") or "", reverse=True)
    return payload[:8]


def _serialize_invoice_ledger(profile):
    invoices = list(
        _tenant_invoice_queryset(profile)
        .prefetch_related("line_items", "payments")
        .order_by("-issue_date", "-due_date", "-id")
    )
    ledger = []
    billed_total = ZERO_DECIMAL
    paid_total = ZERO_DECIMAL
    outstanding_total = ZERO_DECIMAL

    for invoice in invoices:
        payments = list(invoice.payments.all())
        line_items = list(invoice.line_items.all())
        paid_amount = sum((payment.amount for payment in payments), ZERO_DECIMAL)
        balance_due = max((invoice.total_amount or ZERO_DECIMAL) - paid_amount, ZERO_DECIMAL)
        billed_total += invoice.total_amount or ZERO_DECIMAL
        paid_total += paid_amount
        outstanding_total += balance_due
        ledger.append(
            {
                "id": invoice.id,
                "invoice_number": invoice.invoice_number,
                "status": invoice.status,
                "status_display": invoice.get_status_display(),
                "issue_date": invoice.issue_date.isoformat(),
                "due_date": invoice.due_date.isoformat(),
                "total_amount": format(invoice.total_amount or ZERO_DECIMAL, ".2f"),
                "paid_amount": format(paid_amount, ".2f"),
                "balance_due": format(balance_due, ".2f"),
                "notes": invoice.notes or "",
                "line_items": TenantLedgerLineItemSerializer(line_items, many=True).data,
                "payments": TenantLedgerPaymentSerializer(payments, many=True).data,
            }
        )

    overdue_total = sum(
        (
            Decimal(entry["balance_due"])
            for entry in ledger
            if entry["balance_due"] != "0.00" and entry["status"] == Invoice.Status.OVERDUE
        ),
        ZERO_DECIMAL,
    )

    return {
        "summary": {
            "invoice_count": len(ledger),
            "billed_total": format(billed_total, ".2f"),
            "paid_total": format(paid_total, ".2f"),
            "outstanding_total": format(outstanding_total, ".2f"),
            "overdue_total": format(overdue_total, ".2f"),
        },
        "ledger": ledger,
    }


def _major_work_orders_for_profile(profile):
    query = WorkOrder.objects.filter(organization=profile.organization).exclude(status=WorkOrder.Status.CANCELLED)
    filters = []
    if profile.unit_id:
        filters.append({"unit_id": profile.unit_id})
    if profile.facility_space_id:
        filters.append({"facility_space_id": profile.facility_space_id})
    if profile.property_id:
        filters.append({"property_id": profile.property_id})
    if not filters:
        return []

    from django.db.models import Q

    location_filter = Q()
    for condition in filters:
        location_filter |= Q(**condition)

    return list(
        query.filter(location_filter)
        .filter(Q(priority__in=[WorkOrder.Priority.HIGH, WorkOrder.Priority.CRITICAL, WorkOrder.Priority.URGENT]) | Q(is_breakdown=True))
        .order_by("-reported_date", "-id")[:12]
    )


def _serialize_incident_log(profile):
    manual_records = list(
        TenantIncidentRecord.objects.filter(tenant_profile=profile)
        .select_related("work_order")
        .order_by("-occurred_at", "-id")
    )
    maintenance_records = _major_work_orders_for_profile(profile)
    entries = [
        {
            "id": f"incident-{record.id}",
            "entry_type": "incident",
            "title": record.title,
            "incident_type": record.incident_type,
            "incident_type_display": record.get_incident_type_display(),
            "severity": record.severity,
            "severity_display": record.get_severity_display(),
            "status": record.status,
            "status_display": record.get_status_display(),
            "occurred_at": record.occurred_at.isoformat(),
            "resolved_at": record.resolved_at.isoformat() if record.resolved_at else None,
            "description": record.description,
            "notes": record.notes,
            "work_order_id": record.work_order_id,
            "work_order_title": record.work_order.title if record.work_order_id and record.work_order else "",
            "source_label": "Tenant Incident",
        }
        for record in manual_records
    ]

    for work_order in maintenance_records:
        entries.append(
            {
                "id": f"work-order-{work_order.id}",
                "entry_type": "work_order",
                "title": work_order.title,
                "incident_type": TenantIncidentRecord.IncidentType.MAINTENANCE_ISSUE,
                "incident_type_display": "Maintenance Issue",
                "severity": work_order.priority,
                "severity_display": work_order.get_priority_display(),
                "status": work_order.status,
                "status_display": work_order.get_status_display(),
                "occurred_at": work_order.reported_date.isoformat(),
                "resolved_at": work_order.completed_date.isoformat() if work_order.completed_date else None,
                "description": work_order.description,
                "notes": work_order.notes,
                "work_order_id": work_order.id,
                "work_order_title": work_order.title,
                "source_label": "Maintenance Work Order",
            }
        )

    entries.sort(key=lambda item: (item["occurred_at"], item["id"]), reverse=True)
    return {
        "entries": entries[:18],
        "manual_records": TenantIncidentRecordSerializer(manual_records[:12], many=True).data,
        "maintenance_records": [
            {
                "id": work_order.id,
                "title": work_order.title,
                "priority": work_order.priority,
                "priority_display": work_order.get_priority_display(),
                "status": work_order.status,
                "status_display": work_order.get_status_display(),
                "reported_date": work_order.reported_date.isoformat(),
                "completed_date": work_order.completed_date.isoformat() if work_order.completed_date else None,
                "description": work_order.description,
                "notes": work_order.notes,
                "is_breakdown": work_order.is_breakdown,
            }
            for work_order in maintenance_records
        ],
    }


class TenantDashboardOverviewView(APIView):
    permission_classes = [IsAuthenticated, HasRolePermission]
    rbac_sub_module = "tenants.all"
    rbac_action = "view"

    def get(self, request):
        org = _resolve_user_org(request)
        if org is None:
            return Response({"detail": "Organization context required."}, status=400)

        today = timezone.localdate()
        month_start, next_month_start = _month_bounds(today)
        lease_horizon_days = 90
        lease_horizon = today + timedelta(days=lease_horizon_days)

        tenant_profiles = list(
            TenantProfile.objects.filter(organization=org)
            .select_related("customer", "property", "unit", "facility", "facility_space")
            .order_by("lease_end_date", "display_name", "id")
        )
        active_tenants = [
            profile
            for profile in tenant_profiles
            if profile.status == TenantProfile.Status.ACTIVE
        ]

        total_units = Unit.objects.filter(organization=org).count()
        occupied_unit_ids = {profile.unit_id for profile in active_tenants if profile.unit_id}
        occupied_units = len(occupied_unit_ids)
        occupancy_rate = round((occupied_units / total_units) * 100, 2) if total_units else 0.0

        tenant_customer_ids = sorted({profile.customer_id for profile in active_tenants if profile.customer_id})
        invoice_snapshot = _invoice_snapshot(
            organization=org,
            customer_ids=tenant_customer_ids,
            today=today,
            month_start=month_start,
            next_month_start=next_month_start,
        )
        overdue_total = invoice_snapshot["totals"]["overdue_total"]
        overdue_invoice_count = invoice_snapshot["totals"]["overdue_invoice_count"]
        current_month_target_total = invoice_snapshot["totals"]["current_month_target_total"]
        current_month_collections = invoice_snapshot["totals"]["current_month_collections"]
        monthly_revenue_attainment = (
            round((current_month_collections / current_month_target_total) * 100, 2)
            if current_month_target_total
            else None
        )

        open_work_orders = WorkOrder.objects.filter(organization=org).exclude(
            status__in=[
                WorkOrder.Status.COMPLETED,
                WorkOrder.Status.VERIFIED,
                WorkOrder.Status.CANCELLED,
            ]
        )
        open_work_order_count = open_work_orders.count()
        urgent_work_order_count = open_work_orders.filter(
            priority__in=[
                WorkOrder.Priority.HIGH,
                WorkOrder.Priority.CRITICAL,
                WorkOrder.Priority.URGENT,
            ]
        ).count()

        expiring_tenants = [
            profile
            for profile in active_tenants
            if profile.lease_end_date and today <= profile.lease_end_date <= lease_horizon
        ]
        soonest_lease_expiry = min(
            (profile.lease_end_date for profile in expiring_tenants if profile.lease_end_date),
            default=None,
        )

        payload = {
            "generated_at": timezone.now().isoformat(),
            "portfolio_scope": {
                "properties_count": Property.objects.filter(organization=org).count(),
                "active_tenants": len(active_tenants),
                "total_units": total_units,
            },
            "occupancy": {
                "rate": occupancy_rate,
                "occupied_units": occupied_units,
                "total_units": total_units,
            },
            "arrears": {
                "total": format(overdue_total, ".2f"),
                "overdue_invoices": overdue_invoice_count,
                "customers_in_arrears": len(invoice_snapshot["totals"]["customers_in_arrears"]),
            },
            "monthly_revenue": {
                "collections_total": format(current_month_collections, ".2f"),
                "target_total": format(current_month_target_total, ".2f"),
                "attainment_pct": monthly_revenue_attainment,
                "period_start": month_start.isoformat(),
                "period_end": (next_month_start - timedelta(days=1)).isoformat(),
            },
            "maintenance": {
                "open_work_orders": open_work_order_count,
                "urgent_open_work_orders": urgent_work_order_count,
            },
            "lease_expirations": {
                "count": len(expiring_tenants),
                "window_days": lease_horizon_days,
                "soonest_date": soonest_lease_expiry.isoformat() if soonest_lease_expiry else None,
            },
        }
        return Response(payload)


class TenantDashboardLayoutView(APIView):
    permission_classes = [IsAuthenticated, HasRolePermission]
    rbac_sub_module = "tenants.all"
    rbac_action = "view"

    def get(self, request):
        org = _resolve_user_org(request)
        if org is None:
            return Response({"detail": "Organization context required."}, status=400)

        today = timezone.localdate()
        tenant_profiles = list(
            TenantProfile.objects.filter(organization=org)
            .select_related(
                "customer",
                "property",
                "unit",
                "unit__property",
                "facility",
                "facility__property",
                "facility_space",
                "facility_space__unit",
                "facility_space__unit__property",
                "facility_space__facility",
                "facility_space__facility__property",
                "facility_space__occupancy_profile",
            )
            .order_by("display_name", "id")
        )

        tenant_customer_ids = sorted({profile.customer_id for profile in tenant_profiles if profile.customer_id})
        invoice_snapshot = _invoice_snapshot(
            organization=org,
            customer_ids=tenant_customer_ids,
            today=today,
        )
        unit_open_work_orders = _open_work_orders_by_unit(org)

        status_rank = {
            "overdue": 0,
            "notice_given": 1,
            "active": 2,
            "pending_move_in": 3,
            "inactive": 4,
            "moved_out": 5,
        }
        master_list = []
        active_profiles_by_unit = {}

        for profile in tenant_profiles:
            if profile.status == TenantProfile.Status.ACTIVE and profile.unit_id:
                active_profiles_by_unit[profile.unit_id] = profile

            customer_snapshot = invoice_snapshot["by_customer"].get(
                profile.customer_id,
                {"balance_due": ZERO_DECIMAL, "overdue_balance": ZERO_DECIMAL, "overdue_invoices": 0},
            )
            status_key, status_label = _tenant_dashboard_status(profile, customer_snapshot, today)
            debt_status_key, debt_status_label = _tenant_debt_status(customer_snapshot)
            tenancy_type_key, tenancy_type_label = _resolve_profile_tenancy_type(profile)
            notice_date = _profile_notice_date(profile)
            master_list.append(
                {
                    "id": profile.id,
                    "tenant_name": profile.resolved_display_name,
                    "unit_label": _resolve_profile_unit_label(profile),
                    "property_name": _resolve_profile_property_name(profile),
                    "property_location": _resolve_profile_property_location(profile),
                    "lease_expiry": notice_date.isoformat() if notice_date else None,
                    "balance": format(customer_snapshot["balance_due"], ".2f"),
                    "status_key": status_key,
                    "status_label": status_label,
                    "tenancy_type_key": tenancy_type_key,
                    "tenancy_type_label": tenancy_type_label,
                    "debt_status_key": debt_status_key,
                    "debt_status_label": debt_status_label,
                }
            )

        master_list.sort(
            key=lambda row: (
                status_rank.get(row["status_key"], 9),
                row["lease_expiry"] or "9999-12-31",
                row["property_name"],
                row["unit_label"],
                row["tenant_name"],
            )
        )

        occupancy_summary = {
            "occupied": 0,
            "vacant_ready": 0,
            "notice_given": 0,
            "under_maintenance": 0,
        }
        property_groups = {}

        units = (
            Unit.objects.filter(organization=org)
            .select_related(
                "property",
                "facility_space",
                "facility_space__occupancy_profile",
            )
            .order_by("property__name", "floor", "unit_number")
        )

        for unit in units:
            profile = active_profiles_by_unit.get(unit.id)
            facility_space = getattr(unit, "facility_space", None)
            occupancy_profile = getattr(facility_space, "occupancy_profile", None) if facility_space else None
            open_work_orders = unit_open_work_orders.get(unit.id, 0)

            if profile and _is_notice_given(profile, today):
                status_key = "notice_given"
                status_label = "Notice Given"
            elif profile:
                status_key = "occupied"
                status_label = "Occupied"
            elif (occupancy_profile and occupancy_profile.status == "maintenance") or open_work_orders > 0:
                status_key = "under_maintenance"
                status_label = "Under Maintenance"
            else:
                status_key = "vacant_ready"
                status_label = "Vacant / Ready"

            occupancy_summary[status_key] += 1
            property_group = property_groups.setdefault(
                unit.property_id,
                {
                    "property_id": unit.property_id,
                    "property_name": unit.property.name,
                    "totals": {
                        "occupied": 0,
                        "vacant_ready": 0,
                        "notice_given": 0,
                        "under_maintenance": 0,
                    },
                    "units": [],
                },
            )
            property_group["totals"][status_key] += 1

            customer_snapshot = invoice_snapshot["by_customer"].get(
                getattr(profile, "customer_id", None),
                {"balance_due": ZERO_DECIMAL},
            )
            property_group["units"].append(
                {
                    "id": unit.id,
                    "tenant_profile_id": profile.id if profile else None,
                    "unit_number": unit.unit_number,
                    "floor": unit.floor,
                    "space_label": facility_space.space_label if facility_space else "",
                    "location_description": unit.location_description,
                    "tenant_name": profile.resolved_display_name if profile else "",
                    "balance": format(customer_snapshot["balance_due"], ".2f"),
                    "lease_expiry": _profile_notice_date(profile).isoformat()
                    if profile and _profile_notice_date(profile)
                    else None,
                    "status_key": status_key,
                    "status_label": status_label,
                    "open_work_orders": open_work_orders,
                }
            )

        payload = {
            "generated_at": timezone.now().isoformat(),
            "master_list": master_list,
            "filters": {
                "property_locations": sorted(
                    {row["property_location"] for row in master_list if row["property_location"] and row["property_location"] != "Unknown"}
                ),
                "tenancy_types": [
                    {"value": "residential", "label": "Residential"},
                    {"value": "commercial", "label": "Commercial"},
                ],
                "debt_statuses": [
                    {"value": "overdue", "label": "Overdue"},
                    {"value": "outstanding", "label": "Outstanding"},
                    {"value": "clear", "label": "Clear"},
                ],
            },
            "occupancy_vacancy": {
                "totals": occupancy_summary,
                "properties": list(property_groups.values()),
            },
        }
        return Response(payload)


class TenantDashboardRevenueAnalyticsView(APIView):
    permission_classes = [IsAuthenticated, HasRolePermission]
    rbac_sub_module = "tenants.all"
    rbac_action = "view"

    def get(self, request):
        org = _resolve_user_org(request)
        if org is None:
            return Response({"detail": "Organization context required."}, status=400)

        today = timezone.localdate()
        current_month_start, next_month_start = _month_bounds(today)
        month_starts = [_add_months(current_month_start, offset) for offset in range(-5, 1)]
        trend_period_start = month_starts[0]
        trend_period_end = next_month_start
        tenant_customer_ids = _tenant_customer_ids_for_org(org)

        month_totals = {month_start: ZERO_DECIMAL for month_start in month_starts}
        property_income_totals = {}
        payments = (
            InvoicePayment.objects.filter(
                invoice__organization=org,
                invoice__customer_id__in=tenant_customer_ids,
                payment_date__gte=trend_period_start,
                payment_date__lt=trend_period_end,
            )
            .select_related("invoice__property")
            .order_by("payment_date", "id")
        )
        for payment in payments:
            month_start = payment.payment_date.replace(day=1)
            if month_start in month_totals:
                month_totals[month_start] += payment.amount
            property_id = getattr(payment.invoice, "property_id", None)
            if property_id:
                property_income_totals[property_id] = property_income_totals.get(property_id, ZERO_DECIMAL) + payment.amount

        collection_months = [
            {
                "label": _month_label(month_start),
                "month_start": month_start.isoformat(),
                "collections_total": format(month_totals[month_start], ".2f"),
            }
            for month_start in month_starts
        ]
        total_collections = sum((month_totals[month_start] for month_start in month_starts), ZERO_DECIMAL)
        peak_collections = max((month_totals[month_start] for month_start in month_starts), default=ZERO_DECIMAL)
        average_collections = (
            total_collections / Decimal(len(month_starts)) if month_starts else ZERO_DECIMAL
        )

        aging_buckets = {
            "current": {"key": "current", "label": "0-30 Days", "total": ZERO_DECIMAL, "count": 0},
            "overdue_30": {"key": "overdue_30", "label": "31-60 Days", "total": ZERO_DECIMAL, "count": 0},
            "overdue_60": {"key": "overdue_60", "label": "61-90 Days", "total": ZERO_DECIMAL, "count": 0},
            "overdue_90": {"key": "overdue_90", "label": "91+ Days", "total": ZERO_DECIMAL, "count": 0},
        }
        total_overdue = ZERO_DECIMAL
        overdue_invoice_count = 0
        aging_invoices = (
            Invoice.objects.filter(organization=org, customer_id__in=tenant_customer_ids)
            .exclude(status=Invoice.Status.CANCELLED)
            .prefetch_related("payments")
            .order_by("due_date", "id")
        )
        for invoice in aging_invoices:
            paid_amount = sum((payment.amount for payment in invoice.payments.all()), ZERO_DECIMAL)
            balance_due = max(invoice.total_amount - paid_amount, ZERO_DECIMAL)
            if balance_due <= ZERO_DECIMAL or invoice.due_date >= today:
                continue

            days_overdue = (today - invoice.due_date).days
            if days_overdue <= 30:
                bucket = aging_buckets["current"]
            elif days_overdue <= 60:
                bucket = aging_buckets["overdue_30"]
            elif days_overdue <= 90:
                bucket = aging_buckets["overdue_60"]
            else:
                bucket = aging_buckets["overdue_90"]

            bucket["total"] += balance_due
            bucket["count"] += 1
            total_overdue += balance_due
            overdue_invoice_count += 1

        property_cost_totals = {}
        maintenance_work_orders = (
            WorkOrder.objects.filter(
                organization=org,
                completed_date__gte=trend_period_start,
                completed_date__lt=trend_period_end,
                status__in=[WorkOrder.Status.COMPLETED, WorkOrder.Status.VERIFIED],
                actual_cost__isnull=False,
            )
            .select_related("property")
            .order_by("completed_date", "id")
        )
        maintenance_cost_total = ZERO_DECIMAL
        for work_order in maintenance_work_orders:
            actual_cost = work_order.actual_cost or ZERO_DECIMAL
            if actual_cost <= ZERO_DECIMAL:
                continue
            maintenance_cost_total += actual_cost
            if work_order.property_id:
                property_cost_totals[work_order.property_id] = property_cost_totals.get(work_order.property_id, ZERO_DECIMAL) + actual_cost

        property_ids = sorted(set(property_income_totals.keys()) | set(property_cost_totals.keys()))
        property_lookup = {
            record.id: record.name
            for record in Property.objects.filter(organization=org, id__in=property_ids).only("id", "name")
        }
        property_breakdown = []
        for property_id in property_ids:
            income_total = property_income_totals.get(property_id, ZERO_DECIMAL)
            cost_total = property_cost_totals.get(property_id, ZERO_DECIMAL)
            property_breakdown.append(
                {
                    "property_id": property_id,
                    "property_name": property_lookup.get(property_id, f"Property {property_id}"),
                    "income_collected": format(income_total, ".2f"),
                    "maintenance_cost": format(cost_total, ".2f"),
                    "net": format(income_total - cost_total, ".2f"),
                }
            )
        property_breakdown.sort(
            key=lambda row: (
                -(Decimal(row["income_collected"]) + Decimal(row["maintenance_cost"])),
                row["property_name"],
            )
        )

        payload = {
            "generated_at": timezone.now().isoformat(),
            "collection_trend": {
                "months": collection_months,
                "total_collections": format(total_collections, ".2f"),
                "peak_collections": format(peak_collections, ".2f"),
                "average_collections": format(average_collections, ".2f"),
            },
            "aging_buckets": {
                "total_overdue": format(total_overdue, ".2f"),
                "overdue_invoices": overdue_invoice_count,
                "buckets": [
                    {
                        "key": bucket["key"],
                        "label": bucket["label"],
                        "total": format(bucket["total"], ".2f"),
                        "count": bucket["count"],
                    }
                    for bucket in aging_buckets.values()
                ],
            },
            "expense_vs_income": {
                "period_start": trend_period_start.isoformat(),
                "period_end": (trend_period_end - timedelta(days=1)).isoformat(),
                "rent_collected_total": format(total_collections, ".2f"),
                "maintenance_cost_total": format(maintenance_cost_total, ".2f"),
                "net_operating_total": format(total_collections - maintenance_cost_total, ".2f"),
                "expense_ratio_pct": (
                    round((maintenance_cost_total / total_collections) * 100, 2)
                    if total_collections > ZERO_DECIMAL
                    else None
                ),
                "property_breakdown": property_breakdown,
            },
        }
        return Response(payload)


class TenantRegistryOverviewView(APIView):
    permission_classes = [IsAuthenticated, HasRolePermission]
    rbac_sub_module = "tenants.all"
    rbac_action = "view"

    def get(self, request):
        org = _resolve_user_org(request)
        if org is None:
            return Response({"detail": "Organization context required."}, status=400)

        today = timezone.localdate()
        lease_horizon = today + timedelta(days=60)
        move_in_horizon = today + timedelta(days=30)
        recent_move_out_threshold = today - timedelta(days=30)

        profiles = list(
            TenantProfile.objects.filter(organization=org)
            .select_related(
                "customer",
                "contact_account",
                "primary_user",
                "property",
                "unit",
                "unit__property",
                "facility",
                "facility__property",
                "facility_space",
                "facility_space__unit",
                "facility_space__facility",
                "facility_space__facility__property",
            )
            .order_by("display_name", "id")
        )

        total_profiles = len(profiles)
        status_totals = {
            "active": 0,
            "pending_move_in": 0,
            "moved_out": 0,
            "inactive": 0,
        }
        tenant_type_totals = {
            TenantProfile.TenantType.INDIVIDUAL: 0,
            TenantProfile.TenantType.CORPORATE: 0,
        }
        missing_customer_links = 0
        missing_contact_channels = 0
        missing_location_mapping = 0
        unassigned_profiles = 0
        assigned_unit_ids = set()
        assigned_space_ids = set()
        expiring_soon = 0
        upcoming_move_ins = 0
        recent_move_outs = 0

        for profile in profiles:
            status_totals[profile.status] = status_totals.get(profile.status, 0) + 1
            tenant_type_totals[profile.tenant_type] = tenant_type_totals.get(profile.tenant_type, 0) + 1

            if not profile.customer_id:
                missing_customer_links += 1
            if not _tenant_contact_email(profile) and not _tenant_contact_phone(profile):
                missing_contact_channels += 1
            if not any([profile.property_id, profile.unit_id, profile.facility_id, profile.facility_space_id]):
                missing_location_mapping += 1
            if not profile.unit_id and not profile.facility_space_id:
                unassigned_profiles += 1
            if profile.unit_id:
                assigned_unit_ids.add(profile.unit_id)
            if profile.facility_space_id:
                assigned_space_ids.add(profile.facility_space_id)

            lease_watch_date = _profile_notice_date(profile)
            if lease_watch_date and today <= lease_watch_date <= lease_horizon:
                expiring_soon += 1
            if profile.move_in_date and today <= profile.move_in_date <= move_in_horizon:
                upcoming_move_ins += 1
            if profile.move_out_date and recent_move_out_threshold <= profile.move_out_date <= today:
                recent_move_outs += 1

        payload = {
            "generated_at": timezone.now().isoformat(),
            "totals": {
                "profiles": total_profiles,
                "active": status_totals.get(TenantProfile.Status.ACTIVE, 0),
                "pending_move_in": status_totals.get(TenantProfile.Status.PENDING_MOVE_IN, 0),
                "moved_out": status_totals.get(TenantProfile.Status.MOVED_OUT, 0),
                "inactive": status_totals.get(TenantProfile.Status.INACTIVE, 0),
            },
            "hygiene": {
                "missing_customer_links": missing_customer_links,
                "missing_contact_channels": missing_contact_channels,
                "missing_location_mapping": missing_location_mapping,
            },
            "allocation": {
                "assigned_units": len(assigned_unit_ids),
                "assigned_spaces": len(assigned_space_ids),
                "unassigned_profiles": unassigned_profiles,
            },
            "lease_watch": {
                "expiring_next_60_days": expiring_soon,
                "move_ins_next_30_days": upcoming_move_ins,
                "recent_move_outs": recent_move_outs,
            },
            "tenant_types": [
                {
                    "key": TenantProfile.TenantType.CORPORATE,
                    "label": "Corporate",
                    "count": tenant_type_totals.get(TenantProfile.TenantType.CORPORATE, 0),
                },
                {
                    "key": TenantProfile.TenantType.INDIVIDUAL,
                    "label": "Individual",
                    "count": tenant_type_totals.get(TenantProfile.TenantType.INDIVIDUAL, 0),
                },
            ],
        }
        return Response(payload)


class TenantRegistryExportView(APIView):
    permission_classes = [IsAuthenticated, HasRolePermission]
    rbac_sub_module = "tenants.all"
    rbac_action = "view"

    def get(self, request):
        org = _resolve_user_org(request)
        if org is None:
            return Response({"detail": "Organization context required."}, status=400)

        format_key = (request.query_params.get("export_format") or "csv").strip().lower()
        queryset = TenantProfile.objects.filter(organization=org).select_related(
            "customer",
            "contact_account",
            "primary_user",
            "property",
            "unit__property",
            "unit",
            "facility",
            "facility__property",
            "facility_space",
            "facility_space__unit",
            "facility_space__facility",
            "facility_space__facility__property",
        )
        return _tenant_registry_export_response(queryset=queryset, format_key=format_key)


class TenantRegistryLookupView(APIView):
    permission_classes = [IsAuthenticated, HasRolePermission]
    rbac_sub_module = "tenants.all"
    rbac_action = "view"

    def get(self, request):
        org = _resolve_user_org(request)
        if org is None:
            return Response({"detail": "Organization context required."}, status=400)

        customers = [
            {
                "id": customer.id,
                "name": customer.name,
                "email": customer.email,
                "phone": customer.phone,
            }
            for customer in Customer.objects.filter(organization=org).order_by("name")
        ]
        contact_accounts = [
            {
                "id": account.id,
                "name": account.display_name,
                "email": account.email,
                "phone": account.phone or account.secondary_phone,
                "finance_customer_id": account.finance_customer_id,
            }
            for account in ContactAccount.objects.filter(organization=org, is_active=True).order_by("-created_at")
        ]
        users = [
            {
                "id": user.id,
                "name": user.get_full_name().strip() or user.email or user.username,
                "email": user.email,
            }
            for user in User.objects.filter(profile__organization=org).order_by("email").distinct()
        ]
        properties = [
            {
                "id": property_record.id,
                "name": property_record.name,
                "location": _property_location_label(property_record),
                "address": property_record.address,
            }
            for property_record in Property.objects.filter(organization=org).order_by("name")
        ]
        facilities = [
            {
                "id": facility.id,
                "property_id": facility.property_id,
                "property_name": facility.property.name,
                "facility_code": facility.facility_code,
                "label": f"{facility.facility_code} - {facility.property.name}",
            }
            for facility in Facility.objects.filter(organization=org).select_related("property").order_by("property__name")
        ]
        units = [
            {
                "id": unit.id,
                "property_id": unit.property_id,
                "property_name": unit.property.name,
                "unit_number": unit.unit_number,
                "location_description": unit.location_description,
                "status": unit.status,
                "status_display": unit.get_status_display(),
                "label": f"{unit.property.name} - {unit.unit_number}",
            }
            for unit in Unit.objects.filter(organization=org).select_related("property").order_by("property__name", "unit_number")
        ]
        spaces = [
            {
                "id": space.id,
                "facility_id": space.facility_id,
                "facility_code": space.facility.facility_code,
                "property_id": space.facility.property_id,
                "property_name": space.facility.property.name,
                "unit_id": space.unit_id,
                "unit_number": space.unit.unit_number,
                "space_label": space.space_label,
                "label": f"{space.facility.facility_code} - {(space.space_label or space.unit.unit_number)}",
            }
            for space in FacilityUnitSpace.objects.filter(organization=org)
            .select_related("unit", "facility", "facility__property")
            .order_by("facility__property__name", "unit__unit_number")
        ]

        return Response(
            {
                "generated_at": timezone.now().isoformat(),
                "customers": customers,
                "contact_accounts": contact_accounts,
                "users": users,
                "properties": properties,
                "facilities": facilities,
                "units": units,
                "spaces": spaces,
            }
        )


class TenantProfileViewSet(OrgScopedMixin, viewsets.ModelViewSet):
    permission_classes = [IsAuthenticated, HasRolePermission]
    rbac_sub_module = "tenants.all"
    lookup_value_regex = r"\d+"
    rbac_action_map = {
        "list": "view",
        "retrieve": "view",
        "create": "create",
        "update": "edit",
        "partial_update": "edit",
        "destroy": "delete",
        "export_registry": "view",
        "import_registry": "create",
        "source_of_truth": "view",
        "intelligence": "view",
        "generate_invoice": "create",
        "send_payment_reminder": "edit",
        "initiate_eviction_notice": "edit",
        "bulk_service_charge": "create",
        "bulk_holiday_announcement": "edit",
        "create_identity_document": "edit",
        "create_relationship_contact": "edit",
        "create_inventory_item": "edit",
        "create_incident_record": "edit",
    }
    http_method_names = ["get", "post", "put", "patch", "head", "options"]
    search_fields = [
        "display_name",
        "customer__name",
        "contact_account__legal_name",
        "contact_account__first_name",
        "contact_account__last_name",
        "property__name",
        "unit__unit_number",
        "facility__facility_code",
        "facility_space__space_label",
    ]
    ordering_fields = [
        "display_name",
        "status",
        "lease_start_date",
        "lease_end_date",
        "move_in_date",
        "move_out_date",
        "updated_at",
    ]
    ordering = ["display_name", "id"]
    queryset = TenantProfile.objects.all()

    def get_queryset(self):
        return (
            super()
            .get_queryset()
            .select_related(
                "customer",
                "contact_account",
                "primary_user",
                "property",
                "unit__property",
                "unit",
                "facility",
                "facility__property",
                "facility_space",
                "facility_space__unit",
                "facility_space__facility",
                "facility_space__facility__property",
            )
        )

    def get_serializer_class(self):
        if self.action in {"create", "update", "partial_update"}:
            return TenantProfileWriteSerializer
        if self.action == "retrieve":
            return TenantProfileDetailSerializer
        return TenantProfileListSerializer

    def get_serializer_context(self):
        context = super().get_serializer_context()
        if self.action in {"list", "retrieve", "source_of_truth"}:
            org = _resolve_user_org(self.request)
            if org is not None:
                customer_ids = sorted(
                    {
                        customer_id
                        for customer_id in self.get_queryset().values_list("customer_id", flat=True)
                        if customer_id
                    }
                )
                if customer_ids:
                    snapshot = _invoice_snapshot(
                        organization=org,
                        customer_ids=customer_ids,
                        today=timezone.localdate(),
                    )
                    context["tenant_invoice_snapshot_by_customer"] = snapshot["by_customer"]
        return context

    def _registry_import_reference_bundle(self, org):
        return {
            "customers": list(Customer.objects.filter(organization=org).order_by("name")),
            "contact_accounts": list(ContactAccount.objects.filter(organization=org).order_by("-created_at")),
            "users": list(User.objects.filter(profile__organization=org).distinct().order_by("email")),
            "properties": list(Property.objects.filter(organization=org).order_by("name")),
            "facilities": list(Facility.objects.filter(organization=org).select_related("property").order_by("facility_code")),
            "units": list(Unit.objects.filter(organization=org).select_related("property").order_by("unit_number")),
            "spaces": list(
                FacilityUnitSpace.objects.filter(organization=org)
                .select_related("facility", "facility__property", "unit")
                .order_by("space_label", "id")
            ),
        }

    def _resolve_registry_import_related(self, row, *, references):
        row_customer_name = _normalize_lookup_value(row.get("customer_name"))
        row_customer_email = _normalize_lookup_value(row.get("customer_email"))
        row_customer_phone = _normalize_lookup_value(row.get("customer_phone"))
        row_contact_name = _normalize_lookup_value(row.get("contact_account_name"))
        row_contact_email = _normalize_lookup_value(row.get("contact_email"))
        row_contact_phone = _normalize_lookup_value(row.get("contact_phone"))
        row_primary_user_email = _normalize_lookup_value(row.get("primary_user_email"))
        row_property_name = _normalize_lookup_value(row.get("property_name"))
        row_unit_number = _normalize_lookup_value(row.get("unit_number"))
        row_facility_code = _normalize_lookup_value(row.get("facility_code"))
        row_space_label = _normalize_lookup_value(row.get("space_label"))

        customer = next(
            (
                item
                for item in references["customers"]
                if (
                    row_customer_name and _normalize_lookup_value(item.name) == row_customer_name
                )
                or (
                    row_customer_email and _normalize_lookup_value(item.email) == row_customer_email
                )
                or (
                    row_customer_phone and _normalize_lookup_value(item.phone) == row_customer_phone
                )
            ),
            None,
        )
        contact_account = next(
            (
                item
                for item in references["contact_accounts"]
                if (
                    row_contact_name and _normalize_lookup_value(item.display_name) == row_contact_name
                )
                or (
                    row_contact_email and _normalize_lookup_value(item.email) == row_contact_email
                )
                or (
                    row_contact_phone
                    and _normalize_lookup_value(item.phone or item.secondary_phone) == row_contact_phone
                )
            ),
            None,
        )
        primary_user = next(
            (
                item
                for item in references["users"]
                if row_primary_user_email and _normalize_lookup_value(item.email) == row_primary_user_email
            ),
            None,
        )
        property_record = next(
            (
                item
                for item in references["properties"]
                if row_property_name and _normalize_lookup_value(item.name) == row_property_name
            ),
            None,
        )
        facility = next(
            (
                item
                for item in references["facilities"]
                if row_facility_code and _normalize_lookup_value(item.facility_code) == row_facility_code
            ),
            None,
        )

        unit = None
        if row_unit_number:
            candidate_units = [item for item in references["units"] if _normalize_lookup_value(item.unit_number) == row_unit_number]
            if property_record is not None:
                unit = next((item for item in candidate_units if item.property_id == property_record.id), None)
            unit = unit or (candidate_units[0] if candidate_units else None)

        facility_space = None
        if row_space_label:
            candidate_spaces = [
                item for item in references["spaces"] if _normalize_lookup_value(item.space_label) == row_space_label
            ]
            if facility is not None:
                facility_space = next((item for item in candidate_spaces if item.facility_id == facility.id), None)
            facility_space = facility_space or (candidate_spaces[0] if candidate_spaces else None)

        return {
            "customer": customer,
            "contact_account": contact_account,
            "primary_user": primary_user,
            "property": property_record,
            "facility": facility,
            "unit": unit,
            "facility_space": facility_space,
        }

    def _source_of_truth_payload(self, profile):
        today = timezone.localdate()
        invoice_metrics = _tenant_invoice_metrics(profile, today)
        identity_documents = TenantIdentityDocument.objects.filter(tenant_profile=profile).order_by("document_type", "-created_at")
        relationship_contacts = TenantRelationshipContact.objects.filter(tenant_profile=profile).order_by("contact_role", "-is_primary", "full_name")
        inventory_items = TenantInventoryItem.objects.filter(tenant_profile=profile).order_by("item_name", "id")
        inventory_total_quantity = sum(item.quantity or 0 for item in inventory_items)

        return {
            "generated_at": timezone.now().isoformat(),
            "tenant": TenantProfileDetailSerializer(profile, context=self.get_serializer_context()).data,
            "identity": {
                "documents": TenantIdentityDocumentSerializer(identity_documents, many=True).data,
                "next_of_kin": TenantRelationshipContactSerializer(
                    relationship_contacts.filter(contact_role=TenantRelationshipContact.ContactRole.NEXT_OF_KIN),
                    many=True,
                ).data,
                "emergency_contacts": TenantRelationshipContactSerializer(
                    relationship_contacts.filter(contact_role=TenantRelationshipContact.ContactRole.EMERGENCY),
                    many=True,
                ).data,
            },
            "financial": {
                **_serialize_invoice_ledger(profile),
                "current_balance": format(invoice_metrics["current_balance"], ".2f"),
                "overdue_balance": format(invoice_metrics["overdue_balance"], ".2f"),
            },
            "inventory": {
                "items": TenantInventoryItemSerializer(inventory_items, many=True).data,
                "summary": {
                    "item_count": inventory_items.count(),
                    "quantity_total": inventory_total_quantity,
                },
            },
            "incident_log": _serialize_incident_log(profile),
        }

    @action(detail=False, methods=["get"], url_path="export")
    def export_registry(self, request):
        format_key = (request.query_params.get("export_format") or "csv").strip().lower()
        return _tenant_registry_export_response(queryset=self.get_queryset(), format_key=format_key)

    @action(detail=True, methods=["get"], url_path="source-of-truth")
    def source_of_truth(self, request, pk=None):
        profile = self.get_object()
        return Response(self._source_of_truth_payload(profile))

    @action(detail=True, methods=["post"], url_path="identity-documents")
    def create_identity_document(self, request, pk=None):
        profile = self.get_object()
        serializer = TenantIdentityDocumentWriteSerializer(data=request.data, context=self.get_serializer_context())
        serializer.is_valid(raise_exception=True)
        record = serializer.save(organization=profile.organization, tenant_profile=profile)
        return Response(
            {
                "detail": "Identity document added to the tenant registry profile.",
                "record": TenantIdentityDocumentSerializer(record).data,
            },
            status=status.HTTP_201_CREATED,
        )

    @action(detail=True, methods=["post"], url_path="relationship-contacts")
    def create_relationship_contact(self, request, pk=None):
        profile = self.get_object()
        serializer = TenantRelationshipContactWriteSerializer(data=request.data, context=self.get_serializer_context())
        serializer.is_valid(raise_exception=True)
        record = serializer.save(organization=profile.organization, tenant_profile=profile)
        return Response(
            {
                "detail": "Relationship contact added to the tenant registry profile.",
                "record": TenantRelationshipContactSerializer(record).data,
            },
            status=status.HTTP_201_CREATED,
        )

    @action(detail=True, methods=["post"], url_path="inventory-items")
    def create_inventory_item(self, request, pk=None):
        profile = self.get_object()
        serializer = TenantInventoryItemWriteSerializer(data=request.data, context=self.get_serializer_context())
        serializer.is_valid(raise_exception=True)
        record = serializer.save(organization=profile.organization, tenant_profile=profile)
        return Response(
            {
                "detail": "Inventory checklist item added to the tenant profile.",
                "record": TenantInventoryItemSerializer(record).data,
            },
            status=status.HTTP_201_CREATED,
        )

    @action(detail=True, methods=["post"], url_path="incident-records")
    def create_incident_record(self, request, pk=None):
        profile = self.get_object()
        serializer = TenantIncidentRecordWriteSerializer(data=request.data, context=self.get_serializer_context())
        serializer.is_valid(raise_exception=True)
        record = serializer.save(organization=profile.organization, tenant_profile=profile)
        return Response(
            {
                "detail": "Incident record added to the tenant profile.",
                "record": TenantIncidentRecordSerializer(record).data,
            },
            status=status.HTTP_201_CREATED,
        )

    @action(detail=False, methods=["post"], url_path="import")
    def import_registry(self, request):
        upload = request.FILES.get("file")
        if upload is None:
            return Response({"detail": "file is required."}, status=status.HTTP_400_BAD_REQUEST)

        try:
            rows, source_format = _load_registry_import_rows(upload)
        except ValueError as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)
        except Exception as exc:  # pragma: no cover - file parsing guard
            return Response({"detail": f"Could not read the import file: {exc}"}, status=status.HTTP_400_BAD_REQUEST)

        org = _resolve_user_org(request)
        if org is None:
            return Response({"detail": "Organization context required."}, status=status.HTTP_400_BAD_REQUEST)

        references = self._registry_import_reference_bundle(org)
        created_profiles = []
        failed_rows = []

        for index, raw_row in enumerate(rows, start=2):
            row = {str(key or "").strip(): raw_row.get(key, "") for key in raw_row.keys()}
            if not any(str(value or "").strip() for value in row.values()):
                continue

            related = self._resolve_registry_import_related(row, references=references)
            tenant_type = TENANT_TYPE_IMPORT_ALIASES.get(_normalize_lookup_value(row.get("tenant_type")), TenantProfile.TenantType.CORPORATE)
            status_value = TENANT_STATUS_IMPORT_ALIASES.get(
                _normalize_lookup_value(row.get("status")),
                TenantProfile.Status.PENDING_MOVE_IN,
            )

            try:
                occupant_count = int(str(row.get("occupant_count") or "1").strip() or "1")
            except ValueError:
                failed_rows.append(
                    {
                        "row": index,
                        "display_name": str(row.get("display_name") or "").strip() or f"Row {index}",
                        "errors": {"occupant_count": ["Occupant count must be a whole number."]},
                    }
                )
                continue

            payload = {
                "customer": related["customer"].id if related["customer"] else None,
                "contact_account": related["contact_account"].id if related["contact_account"] else None,
                "primary_user": related["primary_user"].id if related["primary_user"] else None,
                "property": related["property"].id if related["property"] else None,
                "unit": related["unit"].id if related["unit"] else None,
                "facility": related["facility"].id if related["facility"] else None,
                "facility_space": related["facility_space"].id if related["facility_space"] else None,
                "tenant_type": tenant_type,
                "status": status_value,
                "display_name": str(row.get("display_name") or "").strip(),
                "lease_start_date": str(row.get("lease_start_date") or "").strip() or None,
                "lease_end_date": str(row.get("lease_end_date") or "").strip() or None,
                "move_in_date": str(row.get("move_in_date") or "").strip() or None,
                "move_out_date": str(row.get("move_out_date") or "").strip() or None,
                "occupant_count": occupant_count,
                "notes": str(row.get("notes") or "").strip(),
            }

            serializer = TenantProfileWriteSerializer(data=payload, context=self.get_serializer_context())
            if not serializer.is_valid():
                failed_rows.append(
                    {
                        "row": index,
                        "display_name": payload["display_name"] or str(row.get("customer_name") or "").strip() or f"Row {index}",
                        "errors": serializer.errors,
                    }
                )
                continue

            with transaction.atomic():
                profile = serializer.save(organization=org)
            created_profiles.append({"id": profile.id, "display_name": profile.resolved_display_name})

        if not created_profiles and failed_rows:
            return Response(
                {
                    "detail": "No tenant rows were imported.",
                    "source_format": source_format,
                    "created_count": 0,
                    "failed_rows": failed_rows,
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        return Response(
            {
                "detail": f"Imported {len(created_profiles)} tenant row{'s' if len(created_profiles) != 1 else ''}.",
                "source_format": source_format,
                "created_count": len(created_profiles),
                "created_profiles": created_profiles,
                "failed_rows": failed_rows,
            }
        )

    def _bulk_profiles_from_request(self, request):
        tenant_ids = request.data.get("tenant_ids")
        if not isinstance(tenant_ids, list) or not tenant_ids:
            return None, Response({"detail": "tenant_ids must be a non-empty list."}, status=status.HTTP_400_BAD_REQUEST)

        profiles = list(
            self.get_queryset()
            .filter(id__in=tenant_ids)
            .select_related(
                "customer",
                "contact_account",
                "primary_user",
                "property",
                "unit",
                "unit__property",
                "facility",
                "facility__property",
                "facility_space",
                "facility_space__unit",
                "facility_space__facility",
                "facility_space__facility__property",
            )
        )
        if not profiles:
            return None, Response({"detail": "No tenant profiles matched the selected IDs."}, status=status.HTTP_404_NOT_FOUND)

        found_ids = {profile.id for profile in profiles}
        missing_ids = [tenant_id for tenant_id in tenant_ids if tenant_id not in found_ids]
        if missing_ids:
            return None, Response(
                {"detail": "Some selected tenants are not available in this organization.", "missing_ids": missing_ids},
                status=status.HTTP_400_BAD_REQUEST,
            )

        return profiles, None

    @action(detail=True, methods=["get"], url_path="intelligence")
    def intelligence(self, request, pk=None):
        profile = self.get_object()
        today = timezone.localdate()
        invoice_metrics = _tenant_invoice_metrics(profile, today)
        status_key, status_label = _tenant_dashboard_status(
            profile,
            {"overdue_invoices": invoice_metrics["overdue_invoices"]},
            today,
        )
        property_name = _resolve_profile_property_name(profile)
        unit_label = _resolve_profile_unit_label(profile)
        contact_email = _tenant_contact_email(profile)
        contact_phone = _tenant_contact_phone(profile)

        payload = {
            "generated_at": timezone.now().isoformat(),
            "tenant": {
                "id": profile.id,
                "name": profile.resolved_display_name,
                "status": status_key,
                "status_label": status_label,
                "tenant_type": profile.tenant_type,
                "tenant_type_display": profile.get_tenant_type_display(),
                "property_name": property_name,
                "unit_label": unit_label,
                "lease_start_date": profile.lease_start_date.isoformat() if profile.lease_start_date else None,
                "lease_end_date": profile.lease_end_date.isoformat() if profile.lease_end_date else None,
                "move_in_date": profile.move_in_date.isoformat() if profile.move_in_date else None,
                "move_out_date": profile.move_out_date.isoformat() if profile.move_out_date else None,
                "current_balance": format(invoice_metrics["current_balance"], ".2f"),
                "overdue_balance": format(invoice_metrics["overdue_balance"], ".2f"),
            },
            "payment_reliability": _payment_reliability_snapshot(invoice_metrics),
            "lease_timeline": _lease_timeline_snapshot(profile, today),
            "communication_log": _serialize_communication_log_entries(profile),
            "contact": {
                "email": contact_email,
                "phone": contact_phone,
            },
            "quick_actions": {
                "can_generate_invoice": bool(profile.customer_id),
                "can_send_payment_reminder": bool(contact_email or contact_phone)
                and invoice_metrics["current_balance"] > ZERO_DECIMAL,
                "can_initiate_eviction_notice": invoice_metrics["overdue_balance"] > ZERO_DECIMAL,
            },
        }
        return Response(payload)

    @action(detail=True, methods=["post"], url_path="generate-invoice")
    def generate_invoice(self, request, pk=None):
        profile = self.get_object()
        if not profile.customer_id or not profile.customer:
            return Response(
                {"detail": "This tenant is not linked to a finance customer record."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        today = timezone.localdate()
        invoice_metrics = _tenant_invoice_metrics(profile, today)
        property_record = _resolve_profile_property(profile)
        unit_label = _resolve_profile_unit_label(profile)
        template_amount = (
            invoice_metrics["latest_invoice_amount"]
            or invoice_metrics["overdue_balance"]
            or invoice_metrics["current_balance"]
            or ZERO_DECIMAL
        )

        with transaction.atomic():
            invoice = Invoice.objects.create(
                organization=profile.organization,
                customer=profile.customer,
                property=property_record,
                status=Invoice.Status.DRAFT,
                issue_date=today,
                due_date=today + timedelta(days=7),
                notes=(
                    f"Generated from tenant intelligence for {profile.resolved_display_name}"
                    f" ({unit_label} / {_resolve_profile_property_name(profile)})."
                ),
            )
            if template_amount > ZERO_DECIMAL:
                InvoiceLineItem.objects.create(
                    invoice=invoice,
                    description=f"Tenant charge template - {unit_label or profile.resolved_display_name}",
                    quantity=Decimal("1.00"),
                    unit_price=template_amount,
                    sort_order=1,
                )
                invoice.recalculate_totals()

            log_entry = _record_tenant_log(
                profile,
                interaction_type=TenantCommunicationLog.InteractionType.INVOICE_GENERATED,
                channel=TenantCommunicationLog.Channel.SYSTEM,
                direction=TenantCommunicationLog.Direction.INTERNAL,
                entry_status=TenantCommunicationLog.Status.RECORDED,
                subject=f"Draft invoice {invoice.invoice_number} generated",
                message=(
                    f"Draft invoice {invoice.invoice_number} created for {profile.resolved_display_name} "
                    f"at {unit_label or '--'}."
                ),
                author=request.user,
                metadata={
                    "invoice_id": invoice.id,
                    "invoice_number": invoice.invoice_number,
                    "invoice_total": format(invoice.total_amount, ".2f"),
                },
            )

        return Response(
            {
                "detail": f"Draft invoice {invoice.invoice_number} generated successfully.",
                "invoice": {
                    "id": invoice.id,
                    "invoice_number": invoice.invoice_number,
                    "status": invoice.status,
                    "status_display": invoice.get_status_display(),
                    "total_amount": format(invoice.total_amount, ".2f"),
                    "due_date": invoice.due_date.isoformat(),
                },
                "communication_log": TenantCommunicationLogSerializer(log_entry).data,
            }
        )

    @action(detail=True, methods=["post"], url_path="send-payment-reminder")
    def send_payment_reminder(self, request, pk=None):
        profile = self.get_object()
        today = timezone.localdate()
        invoice_metrics = _tenant_invoice_metrics(profile, today)
        current_balance = invoice_metrics["current_balance"]

        if current_balance <= ZERO_DECIMAL:
            return Response(
                {"detail": "This tenant does not have an outstanding balance to remind against."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        property_name = _resolve_profile_property_name(profile)
        unit_label = _resolve_profile_unit_label(profile)
        contact_email = _tenant_contact_email(profile)
        contact_phone = _tenant_contact_phone(profile)
        if not contact_email and not contact_phone:
            return Response(
                {"detail": "No tenant email or WhatsApp number is available for this reminder."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        subject = f"Payment reminder - {profile.resolved_display_name}"
        message = (
            f"Dear {profile.resolved_display_name},\n\n"
            f"This is a payment reminder for {unit_label} at {property_name}. "
            f"Your outstanding balance is NGN {format(current_balance, '.2f')}."
        )
        if invoice_metrics["overdue_balance"] > ZERO_DECIMAL:
            message += f" Overdue exposure currently stands at NGN {format(invoice_metrics['overdue_balance'], '.2f')}."
        if invoice_metrics["oldest_overdue_due_date"]:
            message += (
                f" The oldest unpaid amount has been due since "
                f"{invoice_metrics['oldest_overdue_due_date'].isoformat()}."
            )
        message += "\n\nPlease arrange payment or contact the management team for reconciliation.\n"

        communication_logs = []
        successful_channels = []
        channel_errors = []

        if contact_email:
            try:
                sent_count = send_mail(subject, message, None, [contact_email])
                email_status = (
                    TenantCommunicationLog.Status.SENT
                    if sent_count
                    else TenantCommunicationLog.Status.FAILED
                )
                if sent_count:
                    successful_channels.append("Email")
                else:
                    channel_errors.append("Email dispatch did not report a successful send.")
            except Exception as exc:  # pragma: no cover - runtime integration path
                email_status = TenantCommunicationLog.Status.FAILED
                channel_errors.append(f"Email failed: {exc}")

            communication_logs.append(
                _record_tenant_log(
                    profile,
                    interaction_type=TenantCommunicationLog.InteractionType.PAYMENT_REMINDER,
                    channel=TenantCommunicationLog.Channel.EMAIL,
                    direction=TenantCommunicationLog.Direction.OUTBOUND,
                    entry_status=email_status,
                    subject=subject,
                    message=message,
                    author=request.user,
                    metadata={"recipient_email": contact_email},
                )
            )

        if contact_phone:
            try:
                dispatch = get_whatsapp_provider_adapter().send_message(
                    recipient_phone=contact_phone,
                    message=message,
                    metadata={"tenant_profile_id": profile.id, "action": "payment_reminder"},
                )
                if dispatch.status == "failed":
                    whatsapp_status = TenantCommunicationLog.Status.FAILED
                    channel_errors.append("WhatsApp dispatch failed.")
                elif dispatch.status == "queued":
                    whatsapp_status = TenantCommunicationLog.Status.QUEUED
                    successful_channels.append("WhatsApp")
                else:
                    whatsapp_status = TenantCommunicationLog.Status.SENT
                    successful_channels.append("WhatsApp")
                whatsapp_metadata = {
                    "recipient_phone": contact_phone,
                    "provider": dispatch.provider_key,
                    "provider_label": dispatch.provider_label,
                    "provider_message_id": dispatch.provider_message_id,
                    "payload": dispatch.payload,
                }
            except Exception as exc:  # pragma: no cover - runtime integration path
                whatsapp_status = TenantCommunicationLog.Status.FAILED
                whatsapp_metadata = {"recipient_phone": contact_phone, "error": str(exc)}
                channel_errors.append(f"WhatsApp failed: {exc}")

            communication_logs.append(
                _record_tenant_log(
                    profile,
                    interaction_type=TenantCommunicationLog.InteractionType.PAYMENT_REMINDER,
                    channel=TenantCommunicationLog.Channel.WHATSAPP,
                    direction=TenantCommunicationLog.Direction.OUTBOUND,
                    entry_status=whatsapp_status,
                    subject=subject,
                    message=message,
                    author=request.user,
                    metadata=whatsapp_metadata,
                )
            )

        if not successful_channels:
            return Response(
                {
                    "detail": "Payment reminder could not be dispatched on the available channels.",
                    "communication_logs": TenantCommunicationLogSerializer(communication_logs, many=True).data,
                    "errors": channel_errors,
                },
                status=status.HTTP_502_BAD_GATEWAY,
            )

        return Response(
            {
                "detail": f"Payment reminder sent via {', '.join(successful_channels)}.",
                "communication_logs": TenantCommunicationLogSerializer(communication_logs, many=True).data,
            }
        )

    @action(detail=True, methods=["post"], url_path="initiate-eviction-notice")
    def initiate_eviction_notice(self, request, pk=None):
        profile = self.get_object()
        today = timezone.localdate()
        invoice_metrics = _tenant_invoice_metrics(profile, today)
        overdue_balance = invoice_metrics["overdue_balance"]

        if overdue_balance <= ZERO_DECIMAL:
            return Response(
                {"detail": "Eviction notice can only be initiated when the tenant has overdue arrears."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        property_name = _resolve_profile_property_name(profile)
        unit_label = _resolve_profile_unit_label(profile)
        contact_email = _tenant_contact_email(profile)
        notice_date = profile.move_out_date or (today + timedelta(days=30))
        if profile.move_out_date != notice_date:
            profile.move_out_date = notice_date
            profile.save(update_fields=["move_out_date", "updated_at"])

        subject = f"Eviction notice initiated - {profile.resolved_display_name}"
        message = (
            f"Formal eviction recovery has been initiated for {profile.resolved_display_name} "
            f"({unit_label} / {property_name}) due to overdue arrears of "
            f"NGN {format(overdue_balance, '.2f')}. "
            f"The current move-out target is {notice_date.isoformat()}."
        )

        communication_logs = [
            _record_tenant_log(
                profile,
                interaction_type=TenantCommunicationLog.InteractionType.EVICTION_NOTICE,
                channel=TenantCommunicationLog.Channel.SYSTEM,
                direction=TenantCommunicationLog.Direction.INTERNAL,
                entry_status=TenantCommunicationLog.Status.RECORDED,
                subject=subject,
                message=message,
                author=request.user,
                metadata={"notice_date": notice_date.isoformat(), "overdue_balance": format(overdue_balance, ".2f")},
            )
        ]

        if contact_email:
            email_body = (
                f"Dear {profile.resolved_display_name},\n\n"
                f"This is formal notice that recovery action has commenced on your tenancy at "
                f"{property_name} ({unit_label}). The overdue balance currently stands at "
                f"NGN {format(overdue_balance, '.2f')}.\n\n"
                f"Please contact management immediately. The current move-out target has been set to "
                f"{notice_date.isoformat()}.\n"
            )
            try:
                sent_count = send_mail(subject, email_body, None, [contact_email])
                email_status = (
                    TenantCommunicationLog.Status.SENT
                    if sent_count
                    else TenantCommunicationLog.Status.FAILED
                )
            except Exception as exc:  # pragma: no cover - runtime integration path
                email_status = TenantCommunicationLog.Status.FAILED
                email_body = f"{email_body}\n\nDispatch error: {exc}"

            communication_logs.append(
                _record_tenant_log(
                    profile,
                    interaction_type=TenantCommunicationLog.InteractionType.EVICTION_NOTICE,
                    channel=TenantCommunicationLog.Channel.EMAIL,
                    direction=TenantCommunicationLog.Direction.OUTBOUND,
                    entry_status=email_status,
                    subject=subject,
                    message=email_body,
                    author=request.user,
                    metadata={"recipient_email": contact_email, "notice_date": notice_date.isoformat()},
                )
            )

        return Response(
            {
                "detail": f"Eviction notice initiated. Move-out target is {notice_date.isoformat()}.",
                "notice_date": notice_date.isoformat(),
                "communication_logs": TenantCommunicationLogSerializer(communication_logs, many=True).data,
            }
        )

    @action(detail=False, methods=["post"], url_path="bulk-service-charge")
    def bulk_service_charge(self, request):
        profiles, error_response = self._bulk_profiles_from_request(request)
        if error_response is not None:
            return error_response

        try:
            amount = _parse_positive_decimal(request.data.get("amount"), field_name="amount")
        except ValueError as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)

        description = (request.data.get("description") or "").strip() or "Service charge"
        due_in_days_raw = request.data.get("due_in_days", 7)
        try:
            due_in_days = int(due_in_days_raw)
        except (TypeError, ValueError):
            return Response({"detail": "due_in_days must be a whole number."}, status=status.HTTP_400_BAD_REQUEST)
        if due_in_days < 0:
            return Response({"detail": "due_in_days cannot be negative."}, status=status.HTTP_400_BAD_REQUEST)

        notes = (request.data.get("notes") or "").strip()
        send_notifications = bool(request.data.get("send_notifications", True))
        today = timezone.localdate()
        created_invoices = []
        notification_logs = []
        skipped_tenants = []

        for profile in profiles:
            if not profile.customer_id or not profile.customer:
                skipped_tenants.append({"tenant_id": profile.id, "tenant_name": profile.resolved_display_name, "reason": "No finance customer linked."})
                continue

            property_record = _resolve_profile_property(profile)
            unit_label = _resolve_profile_unit_label(profile)
            with transaction.atomic():
                invoice = Invoice.objects.create(
                    organization=profile.organization,
                    customer=profile.customer,
                    property=property_record,
                    status=Invoice.Status.SENT,
                    issue_date=today,
                    due_date=today + timedelta(days=due_in_days),
                    notes=notes or f"Bulk service charge invoice for {profile.resolved_display_name}.",
                )
                InvoiceLineItem.objects.create(
                    invoice=invoice,
                    description=description,
                    quantity=Decimal("1.00"),
                    unit_price=amount,
                    sort_order=1,
                )
                invoice.recalculate_totals()

                _record_tenant_log(
                    profile,
                    interaction_type=TenantCommunicationLog.InteractionType.INVOICE_GENERATED,
                    channel=TenantCommunicationLog.Channel.SYSTEM,
                    direction=TenantCommunicationLog.Direction.INTERNAL,
                    entry_status=TenantCommunicationLog.Status.RECORDED,
                    subject=f"Bulk service charge invoice {invoice.invoice_number} generated",
                    message=(
                        f"{description} invoice {invoice.invoice_number} generated for {profile.resolved_display_name} "
                        f"({unit_label} / {_resolve_profile_property_name(profile)})."
                    ),
                    author=request.user,
                    metadata={
                        "bulk_action": "service_charge_invoice",
                        "invoice_id": invoice.id,
                        "invoice_number": invoice.invoice_number,
                        "invoice_total": format(invoice.total_amount, ".2f"),
                    },
                )

                if send_notifications:
                    contact_email = _tenant_contact_email(profile)
                    if contact_email:
                        subject = f"Service charge invoice {invoice.invoice_number}"
                        message = (
                            f"Dear {profile.resolved_display_name},\n\n"
                            f"A new service charge invoice ({invoice.invoice_number}) has been issued for "
                            f"{unit_label} at {_resolve_profile_property_name(profile)}.\n"
                            f"Amount due: NGN {format(invoice.total_amount, '.2f')}\n"
                            f"Due date: {invoice.due_date.isoformat()}\n"
                        )
                        if notes:
                            message += f"\nNotes: {notes}\n"
                        try:
                            sent_count = send_mail(subject, message, None, [contact_email])
                            email_status = TenantCommunicationLog.Status.SENT if sent_count else TenantCommunicationLog.Status.FAILED
                        except Exception as exc:  # pragma: no cover - runtime integration path
                            email_status = TenantCommunicationLog.Status.FAILED
                            message = f"{message}\nDispatch error: {exc}"

                        notification_logs.append(
                            _record_tenant_log(
                                profile,
                                interaction_type=TenantCommunicationLog.InteractionType.EMAIL,
                                channel=TenantCommunicationLog.Channel.EMAIL,
                                direction=TenantCommunicationLog.Direction.OUTBOUND,
                                entry_status=email_status,
                                subject=subject,
                                message=message,
                                author=request.user,
                                metadata={
                                    "bulk_action": "service_charge_invoice",
                                    "invoice_id": invoice.id,
                                    "recipient_email": contact_email,
                                },
                            )
                        )

            created_invoices.append(
                {
                    "tenant_id": profile.id,
                    "tenant_name": profile.resolved_display_name,
                    "invoice_id": invoice.id,
                    "invoice_number": invoice.invoice_number,
                    "total_amount": format(invoice.total_amount, ".2f"),
                }
            )

        if not created_invoices:
            return Response(
                {"detail": "No service charge invoices were generated.", "skipped": skipped_tenants},
                status=status.HTTP_400_BAD_REQUEST,
            )

        return Response(
            {
                "detail": f"Generated {len(created_invoices)} service charge invoice{'s' if len(created_invoices) != 1 else ''}.",
                "generated_count": len(created_invoices),
                "invoices": created_invoices,
                "notification_count": len(notification_logs),
                "skipped": skipped_tenants,
            }
        )

    @action(detail=False, methods=["post"], url_path="bulk-holiday-announcement")
    def bulk_holiday_announcement(self, request):
        profiles, error_response = self._bulk_profiles_from_request(request)
        if error_response is not None:
            return error_response

        subject = _announcement_subject(request.data.get("subject"))
        message_body = (request.data.get("message") or "").strip()
        if not message_body:
            return Response({"detail": "message is required for a holiday announcement."}, status=status.HTTP_400_BAD_REQUEST)

        delivery_mode = (request.data.get("delivery_mode") or "all_available").strip()
        valid_modes = {"all_available", "email_only", "whatsapp_only"}
        if delivery_mode not in valid_modes:
            return Response({"detail": "delivery_mode must be all_available, email_only, or whatsapp_only."}, status=status.HTTP_400_BAD_REQUEST)

        adapter = None
        if delivery_mode in {"all_available", "whatsapp_only"}:
            adapter = get_whatsapp_provider_adapter()

        delivery_logs = []
        sent_count = 0
        skipped_tenants = []

        for profile in profiles:
            contact_email = _tenant_contact_email(profile)
            contact_phone = _tenant_contact_phone(profile)
            attempted_delivery = False

            if delivery_mode in {"all_available", "email_only"} and contact_email:
                attempted_delivery = True
                try:
                    email_sent = send_mail(subject, message_body, None, [contact_email])
                    email_status = TenantCommunicationLog.Status.SENT if email_sent else TenantCommunicationLog.Status.FAILED
                    if email_sent:
                        sent_count += 1
                except Exception as exc:  # pragma: no cover - runtime integration path
                    email_status = TenantCommunicationLog.Status.FAILED
                    email_message = f"{message_body}\n\nDispatch error: {exc}"
                else:
                    email_message = message_body

                delivery_logs.append(
                    _record_tenant_log(
                        profile,
                        interaction_type=TenantCommunicationLog.InteractionType.EMAIL,
                        channel=TenantCommunicationLog.Channel.EMAIL,
                        direction=TenantCommunicationLog.Direction.OUTBOUND,
                        entry_status=email_status,
                        subject=subject,
                        message=email_message,
                        author=request.user,
                        metadata={"bulk_action": "holiday_announcement", "recipient_email": contact_email},
                    )
                )

            if delivery_mode in {"all_available", "whatsapp_only"} and contact_phone and adapter is not None:
                attempted_delivery = True
                try:
                    dispatch = adapter.send_message(
                        recipient_phone=contact_phone,
                        message=message_body,
                        metadata={"bulk_action": "holiday_announcement", "tenant_profile_id": profile.id},
                    )
                    if dispatch.status == "failed":
                        whatsapp_status = TenantCommunicationLog.Status.FAILED
                    elif dispatch.status == "queued":
                        whatsapp_status = TenantCommunicationLog.Status.QUEUED
                        sent_count += 1
                    else:
                        whatsapp_status = TenantCommunicationLog.Status.SENT
                        sent_count += 1
                    whatsapp_metadata = {
                        "bulk_action": "holiday_announcement",
                        "recipient_phone": contact_phone,
                        "provider": dispatch.provider_key,
                        "provider_message_id": dispatch.provider_message_id,
                        "payload": dispatch.payload,
                    }
                except Exception as exc:  # pragma: no cover - runtime integration path
                    whatsapp_status = TenantCommunicationLog.Status.FAILED
                    whatsapp_metadata = {
                        "bulk_action": "holiday_announcement",
                        "recipient_phone": contact_phone,
                        "error": str(exc),
                    }

                delivery_logs.append(
                    _record_tenant_log(
                        profile,
                        interaction_type=TenantCommunicationLog.InteractionType.WHATSAPP,
                        channel=TenantCommunicationLog.Channel.WHATSAPP,
                        direction=TenantCommunicationLog.Direction.OUTBOUND,
                        entry_status=whatsapp_status,
                        subject=subject,
                        message=message_body,
                        author=request.user,
                        metadata=whatsapp_metadata,
                    )
                )

            if not attempted_delivery:
                skipped_tenants.append(
                    {
                        "tenant_id": profile.id,
                        "tenant_name": profile.resolved_display_name,
                        "reason": "No matching contact channel for this delivery mode.",
                    }
                )

        if not delivery_logs:
            return Response(
                {"detail": "No holiday announcements were delivered.", "skipped": skipped_tenants},
                status=status.HTTP_400_BAD_REQUEST,
            )

        return Response(
            {
                "detail": f"Holiday announcement processed for {len(delivery_logs)} channel dispatch{'es' if len(delivery_logs) != 1 else ''}.",
                "delivery_count": len(delivery_logs),
                "successful_dispatches": sent_count,
                "skipped": skipped_tenants,
            }
        )
